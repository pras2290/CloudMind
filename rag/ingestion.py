import hashlib
import json
import logging
import os
import pdfplumber
import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from fetchers import common
from fetchers.common import (  # noqa: F401 (re-exported for app.py)
    CHROMA_DIR, DOCS_DIR, CHUNK_SIZE, CHUNK_OVERLAP,
    get_embeddings, load_vectorstore, vectorstore_exists
)
from fetchers.cost_fetcher import fetch_all_cost_data
from fetchers.ec2_fetcher import fetch_all_ec2_data
from fetchers.s3_fetcher import fetch_all_s3_data
from fetchers.lambda_fetcher import fetch_all_lambda_data
from fetchers.cloudwatch_fetcher import fetch_all_cloudwatch_data
from fetchers.iam_fetcher import fetch_all_iam_data

logger = logging.getLogger(__name__)

AWS_FETCHERS = [
    ("aws_cost_explorer", fetch_all_cost_data),
    ("aws_ec2", fetch_all_ec2_data),
    ("aws_s3", fetch_all_s3_data),
    ("aws_lambda", fetch_all_lambda_data),
    ("aws_cloudwatch", fetch_all_cloudwatch_data),
    ("aws_iam", fetch_all_iam_data),
]

# Chroma rejects very large single add() calls.
ADD_BATCH_SIZE = 1000


# ── Text splitter ─────────────────────────────────────────────
def get_splitter():
    return RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ".", " ", ""]
    )


# ── File loaders ──────────────────────────────────────────────
def load_pdf(file_path):
    docs = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text(layout=True) or ""
            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={
                        "source": file_path,
                        "page": i + 1,
                        "type": "pdf"
                    }
                ))
    return docs


def load_word(file_path):
    docx = DocxDocument(file_path)
    text = "\n".join([p.text for p in docx.paragraphs if p.text.strip()])
    return [Document(
        page_content=text,
        metadata={"source": file_path, "type": "word"}
    )]


def load_csv(file_path):
    df = pd.read_csv(file_path)
    docs = []
    for i, row in df.iterrows():
        text = "\n".join([f"{col}: {val}" for col, val in row.items()])
        docs.append(Document(
            page_content=text,
            metadata={"source": file_path, "row": i, "type": "csv"}
        ))
    return docs


def load_uploaded_files(upload_path="./uploads"):
    docs = []
    if not os.path.exists(upload_path):
        return docs
    for filename in os.listdir(upload_path):
        file_path = os.path.join(upload_path, filename)
        ext = os.path.splitext(filename)[1].lower()
        try:
            if ext == ".pdf":
                docs.extend(load_pdf(file_path))
            elif ext == ".docx":
                docs.extend(load_word(file_path))
            elif ext == ".csv":
                docs.extend(load_csv(file_path))
            logger.info("Loaded: %s", filename)
        except Exception as e:
            logger.warning("Error loading %s: %s", filename, e)
    return docs


# ── Dedup helpers ─────────────────────────────────────────────
def doc_id(doc):
    """Deterministic ID from chunk content + metadata."""
    key = doc.page_content + "\x00" + json.dumps(
        doc.metadata, sort_keys=True, default=str
    )
    return hashlib.md5(key.encode("utf-8")).hexdigest()


def _is_aws_source(source):
    return isinstance(source, str) and source.startswith("aws_")


def _sync_vectorstore(vectorstore, chunks, failed_sources):
    """Add new chunks; drop stale ones from sources that fetched cleanly."""
    unique = {}
    for chunk in chunks:
        unique.setdefault(doc_id(chunk), chunk)

    existing = vectorstore.get(include=["metadatas"])
    existing_ids = set(existing["ids"])

    stale_ids = []
    for id_, meta in zip(existing["ids"], existing["metadatas"]):
        if id_ in unique:
            continue
        source = (meta or {}).get("source")
        if _is_aws_source(source) and source in failed_sources:
            continue  # keep last good data for services that failed this run
        stale_ids.append(id_)

    new_ids = [i for i in unique if i not in existing_ids]
    logger.info(
        "Chunks: %d unique, %d already indexed, %d new, %d stale removed",
        len(unique), len(unique) - len(new_ids), len(new_ids), len(stale_ids)
    )

    for start in range(0, len(stale_ids), ADD_BATCH_SIZE):
        vectorstore.delete(ids=stale_ids[start:start + ADD_BATCH_SIZE])
    for start in range(0, len(new_ids), ADD_BATCH_SIZE):
        batch = new_ids[start:start + ADD_BATCH_SIZE]
        vectorstore.add_documents([unique[i] for i in batch], ids=batch)


# ── Master ingestion ──────────────────────────────────────────
def ingest_all():
    all_docs = []
    splitter = get_splitter()

    logger.info("── Fetching AWS data ────────────────────")
    common.reset_fetch_failures()
    failed_sources = set()
    for source, fetch in AWS_FETCHERS:
        try:
            all_docs.extend(fetch())
        except Exception:
            logger.exception("Fetching %s failed; continuing", source)
            failed_sources.add(source)
    failed_sources |= common.get_fetch_failures()
    if failed_sources:
        logger.warning("Services with fetch errors: %s", sorted(failed_sources))

    logger.info("── Loading uploaded files ───────────────")
    all_docs.extend(load_uploaded_files())

    logger.info("Total raw documents: %d", len(all_docs))

    # Chunk
    chunks = splitter.split_documents(all_docs)
    logger.info("Total chunks after splitting: %d", len(chunks))

    # Embed + Store (incremental)
    logger.info("Embedding and storing in ChromaDB...")
    vectorstore = load_vectorstore()
    _sync_vectorstore(vectorstore, chunks, failed_sources)
    logger.info("Ingestion complete.")
    return vectorstore


# ── Test ──────────────────────────────────────────────────────
if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    logger.info("Starting full ingestion...")
    vectorstore = ingest_all()
    logger.info("Vectorstore ready — %d chunks indexed.", len(vectorstore.get(include=[])["ids"]))

    # Quick test query
    results = vectorstore.similarity_search("AWS cost", k=3)
    logger.info("Test query results: %d", len(results))
    for r in results:
        logger.info("\n%s", r.page_content[:200])
        logger.info("Source: %s", r.metadata['source'])
